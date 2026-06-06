import streamlit as st
import pandas as pd
import google.generativeai as genai
import json
import io
import os
import urllib.request
from dotenv import load_dotenv
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
import gspread
from google.oauth2.service_account import Credentials
import altair as alt

# Load environment variables
load_dotenv()

# Page configuration
st.set_page_config(
    page_title="Real Estate AI Lead Scoring",
    page_icon="🎯",
    layout="wide",
    initial_sidebar_state="expanded"
)

# Custom premium styling
st.markdown("""
<style>
@import url('https://fonts.googleapis.com/css2?family=Plus+Jakarta+Sans:wght@300;400;500;600;700&display=swap');

html, body, [class*="css"] {
    font-family: 'Plus Jakarta Sans', sans-serif;
}

.main-title {
    background: linear-gradient(135deg, #10B981, #3B82F6);
    -webkit-background-clip: text;
    -webkit-text-fill-color: transparent;
    font-size: 2.3rem;
    font-weight: 700;
    margin-bottom: 0.1rem;
    text-shadow: 0px 4px 20px rgba(16, 185, 129, 0.05);
}

.sub-title {
    color: #6B7280;
    font-size: 1.05rem;
    margin-bottom: 1.5rem;
    font-weight: 400;
}

.kpi-container {
    display: flex;
    gap: 1rem;
    margin-bottom: 1rem;
}

.kpi-card {
    flex: 1;
    background: white;
    border-radius: 16px;
    padding: 1.2rem;
    box-shadow: 0 4px 20px -2px rgba(0, 0, 0, 0.04);
    border: 1px solid #F3F4F6;
    text-align: center;
    transition: transform 0.2s ease, box-shadow 0.2s ease;
}

.kpi-card:hover {
    transform: translateY(-2px);
    box-shadow: 0 10px 25px -5px rgba(0, 0, 0, 0.07);
}

.kpi-card-total { border-top: 5px solid #6B7280; }
.kpi-card-vip { border-top: 5px solid #10B981; }
.kpi-card-normal { border-top: 5px solid #3B82F6; }
.kpi-card-junk { border-top: 5px solid #EF4444; }

.kpi-title {
    font-size: 0.82rem;
    color: #4B5563;
    font-weight: 600;
    text-transform: uppercase;
    letter-spacing: 0.05em;
    margin-bottom: 0.3rem;
}

.kpi-value {
    font-size: 2.1rem;
    font-weight: 700;
    margin: 0;
}

.kpi-value-total { color: #111827; }
.kpi-value-vip { color: #10B981; }
.kpi-value-normal { color: #3B82F6; }
.kpi-value-junk { color: #EF4444; }

/* Dark mode adjustments */
@media (prefers-color-scheme: dark) {
    .kpi-card {
        background: #1F2937;
        border-color: #374151;
        box-shadow: 0 4px 20px -2px rgba(0, 0, 0, 0.25);
    }
    .kpi-card:hover {
        box-shadow: 0 10px 25px -5px rgba(0, 0, 0, 0.35);
    }
    .kpi-title {
        color: #9CA3AF;
    }
    .kpi-value-total {
        color: #F9FAFB;
    }
}
</style>
""", unsafe_allow_html=True)

# Helper function to convert Google Sheet edit link to CSV export link
def get_csv_url(sheet_url):
    if "docs.google.com/spreadsheets" in sheet_url:
        try:
            parts = sheet_url.split("/d/")
            if len(parts) > 1:
                doc_id = parts[1].split("/")[0]
                return f"https://docs.google.com/spreadsheets/d/{doc_id}/export?format=csv"
        except Exception:
            pass
    return sheet_url

# Load Google Sheet using public download or private service account credential
@st.cache_data(show_spinner="Đang kết nối và tải dữ liệu Google Sheets...")
def load_sheet_data_secure(url):
    scope = ['https://spreadsheets.google.com/feeds', 'https://www.googleapis.com/auth/drive']
    creds = None
    
    # 1. Try reading credentials from st.secrets (recommended for Streamlit Cloud)
    try:
        if "gcp_service_account" in st.secrets:
            creds_info = st.secrets["gcp_service_account"]
            if isinstance(creds_info, str):
                creds_info = json.loads(creds_info)
            creds = Credentials.from_service_account_info(creds_info, scopes=scope)
    except Exception as e:
        # Suppress error if secrets file is missing locally
        pass
            
    # 2. Try reading credentials from a local file
    if creds is None and os.path.exists("google_credentials.json"):
        try:
            creds = Credentials.from_service_account_file("google_credentials.json", scopes=scope)
        except Exception as e:
            st.sidebar.warning(f"Lỗi đọc file google_credentials.json: {e}")
            
    # 3. Connect to Private Sheet if Credentials exist
    if creds is not None:
        try:
            client = gspread.authorize(creds)
            if "docs.google.com/spreadsheets" in url:
                parts = url.split("/d/")
                if len(parts) > 1:
                    sheet_id = parts[1].split("/")[0]
                    sheet = client.open_by_key(sheet_id)
                    worksheet = sheet.get_worksheet(0)
                    data = worksheet.get_all_records()
                    
                    df = pd.DataFrame(data)
                    df.columns = [col.strip() for col in df.columns]
                    
                    # Standardize column structure
                    required = ['id', 'ten_khach', 'sdt', 'nhu_cau_mo_ta']
                    missing = [c for c in required if c not in df.columns]
                    if missing:
                        st.error(f"⚠️ Cột bắt buộc bị thiếu trong Sheet: {', '.join(missing)}")
                        return None, None
                        
                    return df, "private"
        except Exception as e:
            st.sidebar.warning(f"⚠️ Không truy cập được dạng Riêng Tư. Lỗi: {e}. Đang thử chế độ Công Khai...")
            
    # 4. Fallback: Public URL Download
    try:
        csv_url = get_csv_url(url)
        req = urllib.request.Request(csv_url, headers={'User-Agent': 'Mozilla/5.0'})
        with urllib.request.urlopen(req) as response:
            csv_data = response.read().decode('utf-8')
            
        df = pd.read_csv(io.StringIO(csv_data))
        
        required = ['id', 'ten_khach', 'sdt', 'nhu_cau_mo_ta']
        missing = [col for col in required if col not in df.columns]
        if missing:
            st.error(f"⚠️ Google Sheet thiếu các cột bắt buộc: {', '.join(missing)}")
            return None, None
            
        return df, "public"
    except Exception as e:
        st.error(f"❌ Không thể tải dữ liệu. Lỗi: {str(e)}")
        return None, None

# Normalize fields helper
def normalize_dataframe(df):
    if df is None:
        return None
    df_copy = df.copy()
    if 'diem' not in df_copy.columns:
        df_copy['diem'] = 0
    else:
        df_copy['diem'] = df_copy['diem'].fillna(0).astype(int)
        
    if 'phan_loai' not in df_copy.columns:
        df_copy['phan_loai'] = "Bình thường"
    else:
        df_copy['phan_loai'] = df_copy['phan_loai'].fillna("Bình thường")
        
    if 'ly_do_chi_tiet' not in df_copy.columns:
        df_copy['ly_do_chi_tiet'] = "Chưa chấm điểm (Chưa chạy AI/Rule)"
    else:
        df_copy['ly_do_chi_tiet'] = df_copy['ly_do_chi_tiet'].fillna("Chưa chấm điểm (Chưa chạy AI/Rule)")
    return df_copy

# Rule-based offline scoring logic
def rule_based_score(description_text):
    desc = str(description_text).lower()
    
    # 1. CHECK JUNK / RÁC RULES (-50 POINTS)
    is_junk = False
    reasons = []
    
    # Unrealistic price
    if (("giá 1-2 tỷ" in desc or "1 tỷ" in desc or "2 tỷ" in desc) and "quận 1" in desc) or \
       ("trung tâm" in desc and "vài trăm triệu" in desc) or \
       ("thuê nguyên căn" in desc and "2 triệu" in desc) or \
       ("trung tâm thành phố" in desc and "2 triệu" in desc):
        is_junk = True
        reasons.append("Yêu cầu phi thực tế về giá cả")
        
    # No demand
    if "nhầm số" in desc or "không có nhu cầu" in desc or "dữ liệu cũ" in desc or "nhầm ngành" in desc:
        is_junk = True
        reasons.append("Khách hàng không có nhu cầu hoặc nhầm số/dữ liệu cũ")
        
    # No goodwill
    if "hỏi giá cho vui" in desc or "chưa có ý định mua" in desc or "thái độ không hợp tác" in desc:
        is_junk = True
        reasons.append("Khách hàng không thiện chí/hỏi giá cho vui")
        
    # Spam/Advertising
    if "spam" in desc or "quảng cáo" in desc or "bảo hiểm" in desc or "mời chào" in desc or "vay vốn" in desc:
        is_junk = True
        reasons.append("Nội dung spam hoặc quảng cáo ngược")
        
    # Contact error
    if "thuê bao" in desc or "không bắt máy" in desc or "không phản hồi" in desc or "gọi nhiều lần" in desc:
        is_junk = True
        reasons.append("Thông tin liên lạc lỗi (thuê bao/không nghe máy)")
        
    if is_junk:
        return -50, "Rác", " | ".join(reasons)
        
    # 2. CHECK VIP / SIÊU TIỀM NĂNG RULES (+50 POINTS)
    is_vip = False
    vip_reasons = []
    
    # Check if this is a mid-range case that should be overridden to Normal (0 points)
    is_mid_range = False
    if "đất nền vùng ven" in desc or "long an" in desc or "đồng nai" in desc or "2-3 tỷ" in desc:
        is_mid_range = True
    if "căn hộ 2pn" in desc or "quận 7" in desc or "4-5 tỷ" in desc or "gia đình trẻ" in desc:
        is_mid_range = True
    if "thuê mặt bằng" in desc and "dưới 50 triệu" in desc:
        is_mid_range = True
        
    if not is_mid_range:
        # Large budget
        if "tài chính cực mạnh" in desc or "tài chính mạnh" in desc or "không thành vấn đề" in desc or \
           "trên 30 tỷ" in desc or "thanh toán thẳng" in desc or "20 tỷ" in desc or "ngân sách lớn" in desc:
            is_vip = True
            vip_reasons.append("Tài chính lớn/cực mạnh")
            
        # Premium property type
        if "biệt thự đơn lập" in desc or "penthouse" in desc or "shophouse mặt đường lớn" in desc or \
           "quỹ đất công nghiệp" in desc or "sàn văn phòng diện tích lớn" in desc or "trên 2000m2" in desc:
            is_vip = True
            vip_reasons.append("Quan tâm sản phẩm bất động sản cao cấp")
            
        # Prime location
        if ("vinhomes ocean park" in desc or "phú mỹ hưng" in desc or "ven sông" in desc or "khu đông" in desc or "quận 1" in desc) and \
           ("vip" in desc or "biệt thự" in desc or "shophouse" in desc or "doanh nghiệp" in desc or "tài chính" in desc or "sỉ" in desc or "sàn văn phòng" in desc):
            is_vip = True
            vip_reasons.append("Vị trí đắc địa / Khu vực cao cấp")
            
        # Target client
        if "chủ doanh nghiệp" in desc or "nhà đầu tư chuyên nghiệp" in desc or "mua sỉ" in desc or "gom sỉ" in desc or "số lượng lớn" in desc:
            is_vip = True
            vip_reasons.append("Khách hàng là Chủ doanh nghiệp / Nhà đầu tư sỉ")
            
        # Urgency & Transparency
        if "pháp lý chuẩn 100%" in desc or "sổ hồng riêng" in desc or "gặp trực tiếp chủ đầu tư" in desc or "gặp trực tiếp giám đốc" in desc:
            is_vip = True
            vip_reasons.append("Yêu cầu pháp lý cao và giao dịch trực tiếp")
            
    if is_vip:
        return 50, "VIP", " | ".join(vip_reasons)
        
    # 3. DEFAULT: NORMAL / TIỀM NĂNG TRUNG BÌNH (0 POINTS)
    normal_reasons = []
    if "đất nền" in desc or "long an" in desc or "đồng nai" in desc:
        normal_reasons.append("Khách mua đất nền vùng ven tầm trung")
    elif "căn hộ" in desc or "chung cư" in desc:
        normal_reasons.append("Khách mua chung cư tầm trung")
    elif "thuê mặt bằng" in desc:
        normal_reasons.append("Khách thuê mặt bằng kinh doanh")
    else:
        normal_reasons.append("Khách hàng có nhu cầu bất động sản thông thường")
        
    if "vay ngân hàng" in desc or "vay" in desc:
        normal_reasons.append("cần hỗ trợ tài chính/vay ngân hàng")
        
    return 0, "Bình thường", " - ".join(normal_reasons)

# Helper to score a single lead using Gemini API
def score_single_lead(model, client_name, desc, prompt_template):
    prompt = prompt_template.format(ten_khach=client_name, nhu_cau_mo_ta=desc)
    try:
        response = model.generate_content(
            prompt,
            generation_config={"response_mime_type": "application/json"}
        )
        data = json.loads(response.text)
        score = int(data.get("diem", 0))
        classification = str(data.get("phan_loai", "Bình thường"))
        reason = str(data.get("ly_do", "Không có lý do chi tiết."))
        return score, classification, reason
    except Exception as e:
        return 0, "Bình thường", f"Lỗi gọi API: {str(e)}"

# Helper to style and export to Excel using openpyxl
def export_to_excel(df):
    output = io.BytesIO()
    wb = Workbook()
    ws = wb.active
    ws.title = "Lead Scoring Reports"
    ws.views.sheetView[0].showGridLines = True
    
    headers = ["ID", "Tên Khách Hàng", "Số Điện Thoại", "Mô Tả Nhu Cầu", "Điểm Số", "Phân Loại", "Lý Do Chi Tiết"]
    ws.append(headers)
    
    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    thin_border_side = Side(border_style="thin", color="D3D3D3")
    cell_border = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)
    
    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = cell_border
        
    ws.row_dimensions[1].height = 28
    
    vip_fill = PatternFill(start_color="D1E7DD", end_color="D1E7DD", fill_type="solid")
    vip_font = Font(name="Segoe UI", size=10, color="0F5132", bold=True)
    
    normal_fill = PatternFill(start_color="E2F0D9", end_color="E2F0D9", fill_type="solid")
    normal_font = Font(name="Segoe UI", size=10, color="385723", bold=True)
    
    junk_fill = PatternFill(start_color="F8D7DA", end_color="F8D7DA", fill_type="solid")
    junk_font = Font(name="Segoe UI", size=10, color="842029", bold=True)
    
    for r_idx, row in enumerate(df.itertuples(index=False), start=2):
        ws.append([
            row.id,
            row.ten_khach,
            str(row.sdt),
            row.nhu_cau_mo_ta,
            row.diem,
            row.phan_loai,
            row.ly_do_chi_tiet
        ])
        ws.row_dimensions[r_idx].height = 24
        
        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=r_idx, column=col_idx)
            cell.border = cell_border
            cell.font = Font(name="Segoe UI", size=10)
            
            if col_idx in [1, 3, 5, 6]:
                cell.alignment = Alignment(horizontal="center", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
                
            if col_idx == 6:
                val = str(cell.value)
                if val == "VIP":
                    cell.fill = vip_fill
                    cell.font = vip_font
                elif val == "Rác":
                    cell.fill = junk_fill
                    cell.font = junk_font
                else:
                    cell.fill = normal_fill
                    cell.font = normal_font
                    
    column_widths = {1: 8, 2: 22, 3: 15, 4: 55, 5: 10, 6: 15, 7: 45}
    for col_idx, width in column_widths.items():
        col_letter = get_column_letter(col_idx)
        ws.column_dimensions[col_letter].width = width
        
    wb.save(output)
    output.seek(0)
    return output.getvalue()

# SIDEBAR CONFIGURATION
st.sidebar.markdown("### ⚙️ Cấu Hình Hệ Thống")

default_sheet = "https://docs.google.com/spreadsheets/d/1joAy1H6PU19kwgsn57CSk_8cdci6vcDmME21CV_4n6E/edit?usp=sharing"
sheet_url = st.sidebar.text_input(
    "URL Google Sheets:",
    value=default_sheet,
    help="Hỗ trợ link công khai hoặc riêng tư (khi cấu hình Service Account)."
)

gemini_key = st.sidebar.text_input(
    "Gemini API Key (Không bắt buộc):",
    value=os.getenv("GEMINI_API_KEY", ""),
    type="password",
    help="Nếu bỏ trống, hệ thống sẽ tự động sử dụng thuật toán quy tắc offline."
)

# Connect to google sheets and cache
if "master_df" not in st.session_state:
    df_raw, connection_type = load_sheet_data_secure(sheet_url)
    if df_raw is not None:
        df_norm = normalize_dataframe(df_raw)
        st.session_state.master_df = df_norm.copy()
        st.session_state.original_df = df_norm.copy()
        st.session_state.connection_status = connection_type
    else:
        st.session_state.master_df = None
        st.session_state.original_df = None
        st.session_state.connection_status = None

# Refresh button
if st.sidebar.button("🔄 Đồng bộ Google Sheets", use_container_width=True):
    st.cache_data.clear()
    df_raw, connection_type = load_sheet_data_secure(sheet_url)
    if df_raw is not None:
        df_norm = normalize_dataframe(df_raw)
        st.session_state.master_df = df_norm.copy()
        st.session_state.original_df = df_norm.copy()
        st.session_state.connection_status = connection_type
        st.success("Đồng bộ thành công!")
        st.rerun()

# Display connection status badge
if st.session_state.connection_status == "private":
    st.sidebar.success("🟢 Trạng thái: Kết nối Riêng Tư (Secure)")
elif st.session_state.connection_status == "public":
    st.sidebar.info("🔵 Trạng thái: Chế độ Công Khai (Public)")
else:
    st.sidebar.error("🔴 Trạng thái: Chưa có dữ liệu")

# AI System Prompt Customizer
st.sidebar.markdown("### 📝 AI System Prompt Customizer")
default_prompt = """Bạn là chuyên gia phân tích dữ liệu và chấm điểm khách hàng tiềm năng (Lead Scoring) trong ngành Bất Động Sản.
Nhiệm vụ của bạn là phân tích đoạn mô tả nhu cầu khách hàng dưới đây và đưa ra đánh giá dựa trên bộ quy tắc chính xác sau:

1. QUY TẮC CỘNG 50 ĐIỂM (Khách hàng VIP/Siêu tiềm năng - Phân loại: "VIP"):
- Ngân sách lớn: Có đề cập đến số tiền từ 20 tỷ trở lên hoặc các cụm từ "tài chính mạnh", "ngân sách không thành vấn đề", "ngân sách trên 30 tỷ", "thanh toán thẳng".
- Loại hình cao cấp: Tìm kiếm "Biệt thự đơn lập", "Penthouse", "Shophouse mặt đường lớn", "Quỹ đất công nghiệp", "Sàn văn phòng diện tích lớn" (hoặc diện tích > 2000m2).
- Vị trí đắc địa: Yêu cầu các khu vực như "Quận 1", "Ven sông", "Vinhomes Ocean Park", "Phú Mỹ Hưng", "khu Đông".
- Đối tượng khách hàng: Đề cập là "Chủ doanh nghiệp", "Nhà đầu tư chuyên nghiệp", "Mua sỉ", "Mua số lượng lớn", "gom sỉ 5-10 căn".
- Tính cấp thiết & Minh bạch: Yêu cầu "Pháp lý chuẩn 100%", "Sổ hồng riêng", "Muốn gặp trực tiếp chủ đầu tư để đàm phán", "cần gặp trực tiếp giám đốc dự án".

2. QUY TẮC TRỪ 50 ĐIỂM (Khách hàng Rác/Không tiềm năng - Phân loại: "Rác"):
- Yêu cầu phi thực tế: Tìm mua bất động sản với giá thấp vô lý so với thị trường (VD: Nhà Quận 1 giá 1-2 tỷ, nhà trung tâm có sân vườn hồ bơi giá vài trăm triệu, tìm nhà thuê nguyên căn giá 2 triệu ở trung tâm thành phố).
- Không có nhu cầu: "Nhầm số", "Không có nhu cầu", "Dữ liệu cũ", "Nhầm ngành".
- Khách hàng không thiện chí: "Hỏi giá cho vui", "Chưa có ý định mua", "Thái độ không hợp tác".
- Spam/Quảng cáo: Nội dung chứa các dịch vụ khác như "Bảo hiểm", "Vay vốn", "Mời chào dịch vụ", "quảng cáo ngược lại dịch vụ bảo hiểm".
- Thông tin liên lạc lỗi: "Thuê bao", "Gọi nhiều lần không bắt máy", "Không phản hồi Zalo".

3. CÁC TRƯỜNG HỢP KHÁC (Giữ nguyên điểm hoặc cộng ít - Phân loại: "Bình thường" - Điểm: 0):
- Khách hàng tìm mua chung cư, nhà phố tầm trung (3-10 tỷ).
- Khách hàng cần vay ngân hàng, đang cân nhắc chính sách (ví dụ: cần hỗ trợ vay ngân hàng 70%).
- Khách hàng có nhu cầu thực nhưng cần tư vấn thêm về pháp lý hoặc vị trí.
- Thuê mặt bằng kinh doanh spa Quận 1 diện tích khoảng 80-100m2, giá thuê mong muốn dưới 50 triệu/tháng.
- Cần mua đất nền vùng ven (Long An, Đồng Nai) để đầu tư dài hạn với tài chính 2-3 tỷ.

Thông tin khách hàng cần phân tích:
- Tên khách: {ten_khach}
- Mô tả nhu cầu: {nhu_cau_mo_ta}

Hãy trả về kết quả dưới dạng JSON duy nhất với cấu trúc sau:
{{
  "diem": [Điểm số: 50, 0, hoặc -50],
  "phan_loai": "[VIP / Bình thường / Rác]",
  "ly_do": "[Giải thích chi tiết lý do chấm điểm và nhận diện từ khóa/ngữ cảnh]"
}}"""

system_prompt = st.sidebar.text_area(
    "AI Prompt Template:",
    value=default_prompt,
    height=200,
    help="Dùng cho mô hình AI khi có Gemini API Key."
)

# Run scoring button
if st.sidebar.button("⚡ Chạy thuật toán chấm điểm (Lead Scoring)", use_container_width=True):
    if st.session_state.master_df is None:
        st.sidebar.error("⚠️ Không có dữ liệu khách hàng để chấm điểm!")
    else:
        df_to_score = st.session_state.master_df.copy()
        total_leads = len(df_to_score)
        
        progress_bar = st.sidebar.progress(0)
        status_text = st.sidebar.empty()
        
        if not gemini_key:
            # Running offline matching
            status_text.text("⚡ Đang chạy thuật toán Rule-Based Offline...")
            for i in range(total_leads):
                row = df_to_score.iloc[i]
                score, classification, reason = rule_based_score(row['nhu_cau_mo_ta'])
                df_to_score.at[i, 'diem'] = score
                df_to_score.at[i, 'phan_loai'] = classification
                df_to_score.at[i, 'ly_do_chi_tiet'] = reason
                progress_bar.progress((i + 1) / total_leads)
                
            st.session_state.master_df = df_to_score.copy()
            status_text.text("🎉 Đã hoàn tất chấm điểm offline!")
            st.info("ℹ️ Chấm điểm offline bằng quy tắc nghiệp vụ thành công!")
            st.rerun()
        else:
            # Running online Gemini
            try:
                genai.configure(api_key=gemini_key)
                model = genai.GenerativeModel("gemini-1.5-flash")
                
                for i in range(total_leads):
                    row = df_to_score.iloc[i]
                    status_text.text(f"Đang phân tích AI {i+1}/{total_leads}: {row['ten_khach']}")
                    score, classification, reason = score_single_lead(
                        model,
                        row['ten_khach'],
                        row['nhu_cau_mo_ta'],
                        system_prompt
                    )
                    df_to_score.at[i, 'diem'] = score
                    df_to_score.at[i, 'phan_loai'] = classification
                    df_to_score.at[i, 'ly_do_chi_tiet'] = reason
                    progress_bar.progress((i + 1) / total_leads)
                    
                st.session_state.master_df = df_to_score.copy()
                status_text.text("🎉 Đã hoàn tất chấm điểm bằng AI!")
                st.success("Chấm điểm thành công bằng mô hình Gemini!")
                st.rerun()
            except Exception as e:
                st.sidebar.error(f"Lỗi gọi AI: {str(e)}")

# MAIN INTERFACE
# Render Branding Banner
if os.path.exists("real_estate_ai_banner.png"):
    st.image("real_estate_ai_banner.png", use_container_width=True)

st.markdown('<div class="main-title">🎯 AI Lead Scoring & Automation Dashboard</div>', unsafe_allow_html=True)
st.markdown('<div class="sub-title">Bảo mật truy xuất nguồn dữ liệu và giao diện báo cáo chuyên sâu</div>', unsafe_allow_html=True)

if st.session_state.master_df is not None:
    df_stats = st.session_state.master_df
    total_count = len(df_stats)
    vip_count = len(df_stats[df_stats['phan_loai'] == 'VIP'])
    normal_count = len(df_stats[df_stats['phan_loai'] == 'Bình thường'])
    junk_count = len(df_stats[df_stats['phan_loai'] == 'Rác'])
    
    # Grid: KPIs left, Chart right
    col_kpi, col_chart = st.columns([3, 2])
    
    with col_kpi:
        # Render premium metrics cards
        st.markdown(f"""
        <div class="kpi-container">
            <div class="kpi-card kpi-card-total">
                <div class="kpi-title">Tổng số lead</div>
                <div class="kpi-value kpi-value-total">{total_count}</div>
            </div>
            <div class="kpi-card kpi-card-vip">
                <div class="kpi-title">Khách VIP</div>
                <div class="kpi-value kpi-value-vip">{vip_count}</div>
            </div>
        </div>
        <div class="kpi-container">
            <div class="kpi-card kpi-card-normal">
                <div class="kpi-title">Bình thường</div>
                <div class="kpi-value kpi-value-normal">{normal_count}</div>
            </div>
            <div class="kpi-card kpi-card-junk">
                <div class="kpi-title">Khách Rác</div>
                <div class="kpi-value kpi-value-junk">{junk_count}</div>
            </div>
        </div>
        """, unsafe_allow_html=True)
        
    with col_chart:
        # Create horizontal bar chart using Altair
        chart_data = pd.DataFrame({
            'Phân loại': ['VIP', 'Bình thường', 'Rác'],
            'Số lượng': [vip_count, normal_count, junk_count]
        })
        
        chart = alt.Chart(chart_data).mark_bar(cornerRadiusEnd=8, size=24).encode(
            x=alt.X('Số lượng:Q', title='Số lượng lead'),
            y=alt.Y('Phân loại:N', sort=['VIP', 'Bình thường', 'Rác'], title=None),
            color=alt.Color('Phân loại:N', scale=alt.Scale(
                domain=['VIP', 'Bình thường', 'Rác'],
                range=['#10B981', '#3B82F6', '#EF4444']
            ), legend=None)
        ).properties(
            height=160,
            title="Biểu đồ phân loại khách hàng"
        ).configure_title(
            fontSize=14,
            font='Plus Jakarta Sans',
            anchor='start',
            color='#4B5563'
        )
        
        st.altair_chart(chart, use_container_width=True)
        
    # Filters Section
    st.markdown("### 🔍 Bộ lọc & Kiểm duyệt kết quả (Human-In-The-Loop)")
    col_search, col_filter, col_reset = st.columns([2, 1, 1])
    
    with col_search:
        search_query = st.text_input("Tìm kiếm theo Tên hoặc Mô tả nhu cầu:", "")
        
    with col_filter:
        filter_class = st.selectbox("Lọc phân loại khách hàng:", ["Tất cả", "VIP", "Bình thường", "Rác"])
        
    with col_reset:
        st.write(" ")
        st.write(" ")
        if st.button("🔄 Khôi phục dữ liệu ban đầu", use_container_width=True):
            st.session_state.master_df = st.session_state.original_df.copy()
            st.success("Đã khôi phục dữ liệu gốc!")
            st.rerun()
            
    # Filter the displayed data
    display_df = st.session_state.master_df.copy()
    if search_query:
        display_df = display_df[
            display_df['ten_khach'].str.contains(search_query, case=False, na=False) |
            display_df['nhu_cau_mo_ta'].str.contains(search_query, case=False, na=False)
        ]
    if filter_class != "Tất cả":
        display_df = display_df[display_df['phan_loai'] == filter_class]
        
    # Data editor for human review
    st.info("💡 Mẹo: Bạn có thể nhấp đôi chuột vào ô **Điểm số**, **Phân loại** hoặc **Lý do chi tiết (Ghi chú)** bên dưới để chỉnh sửa trực tiếp. Điểm số và phân loại sẽ tự động đồng bộ hóa.")
    
    # Configure columns
    edited_display_df = st.data_editor(
        display_df,
        column_config={
            "id": st.column_config.NumberColumn("ID", disabled=True),
            "ten_khach": st.column_config.TextColumn("Họ & Tên", disabled=True),
            "sdt": st.column_config.TextColumn("Số Điện Thoại", disabled=True),
            "nhu_cau_mo_ta": st.column_config.TextColumn("Mô tả nhu cầu", disabled=True, width="large"),
            "diem": st.column_config.NumberColumn("Điểm số", min_value=-50, max_value=50, step=50),
            "phan_loai": st.column_config.SelectboxColumn(
                "Phân loại",
                options=["VIP", "Bình thường", "Rác"],
                required=True
            ),
            "ly_do_chi_tiet": st.column_config.TextColumn("Lý do chi tiết (Ghi chú)", width="large")
        },
        use_container_width=True,
        num_rows="fixed",
        key="leads_editor_key"
    )
    
    # Synchronize edits back to st.session_state.master_df using relative index to unique 'id'
    if "leads_editor_key" in st.session_state:
        edits = st.session_state.leads_editor_key
        if edits and "edited_rows" in edits and edits["edited_rows"]:
            has_changes = False
            for rel_idx_str, changes in edits["edited_rows"].items():
                rel_idx = int(rel_idx_str)
                if rel_idx < len(display_df):
                    actual_id = display_df.iloc[rel_idx]['id']
                    
                    for col, val in changes.items():
                        # Automatic classification sync if score is changed
                        if col == 'diem':
                            val = int(val)
                            if val >= 50:
                                st.session_state.master_df.loc[st.session_state.master_df['id'] == actual_id, 'phan_loai'] = "VIP"
                            elif val <= -50:
                                st.session_state.master_df.loc[st.session_state.master_df['id'] == actual_id, 'phan_loai'] = "Rác"
                            else:
                                st.session_state.master_df.loc[st.session_state.master_df['id'] == actual_id, 'phan_loai'] = "Bình thường"
                                
                        # Automatic score sync if classification is changed
                        if col == 'phan_loai':
                            if val == "VIP":
                                st.session_state.master_df.loc[st.session_state.master_df['id'] == actual_id, 'diem'] = 50
                            elif val == "Rác":
                                st.session_state.master_df.loc[st.session_state.master_df['id'] == actual_id, 'diem'] = -50
                            else:
                                st.session_state.master_df.loc[st.session_state.master_df['id'] == actual_id, 'diem'] = 0
                                
                        st.session_state.master_df.loc[st.session_state.master_df['id'] == actual_id, col] = val
                        has_changes = True
            if has_changes:
                st.rerun()
                
    # Download section
    st.markdown("### 📥 Kết xuất dữ liệu báo cáo")
    excel_data = export_to_excel(st.session_state.master_df)
    
    st.download_button(
        label="📥 Tải Báo Cáo Excel Kết Quả Chấm Điểm (Bàn Giao)",
        data=excel_data,
        file_name="bao_cao_lead_scoring_bat_dong_san.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        use_container_width=True,
        help="Xuất toàn bộ danh sách đã chấm điểm và kiểm duyệt thành file Excel có định dạng chuyên nghiệp."
    )
else:
    st.warning("⚠️ Không thể kết nối và tải dữ liệu từ Google Sheets. Vui lòng kiểm tra cấu hình trong sidebar.")
