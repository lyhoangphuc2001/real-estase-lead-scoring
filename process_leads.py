import pandas as pd
import io
import urllib.request
import os
import json
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
import gspread
from google.oauth2.service_account import Credentials

def get_csv_url(sheet_url):
    if "docs.google.com/spreadsheets" in sheet_url:
        try:
            parts = sheet_url.split("/d/")
            if len(parts) > 1:
                doc_id = parts[1].split("/")[0]
                return f"https://docs.google.com/spreadsheets/d/{doc_id}/export?format=csv"
        except Exception:
            pass
    return sheet_url

def load_sheet_data_secure(url):
    scope = ['https://spreadsheets.google.com/feeds', 'https://www.googleapis.com/auth/drive']
    creds = None
    
    # 1. Try reading credentials from a local file
    if os.path.exists("google_credentials.json"):
        try:
            print("🔑 Tìm thấy file google_credentials.json, đang kết nối Google Sheets API...")
            creds = Credentials.from_service_account_file("google_credentials.json", scopes=scope)
            client = gspread.authorize(creds)
            
            if "docs.google.com/spreadsheets" in url:
                parts = url.split("/d/")
                if len(parts) > 1:
                    sheet_id = parts[1].split("/")[0]
                    sheet = client.open_by_key(sheet_id)
                    worksheet = sheet.get_worksheet(0)
                    data = worksheet.get_all_records()
                    
                    df = pd.DataFrame(data)
                    df.columns = [col.strip() for col in df.columns]
                    print("🔒 Kết nối thành công Google Sheet Riêng Tư bằng Service Account!")
                    return df
        except Exception as e:
            print(f"⚠️ Cảnh báo: Lỗi kết nối Riêng Tư: {e}. Đang chuyển sang đọc Công Khai...")
            
    # 2. Fallback: Public URL Download
    print("🌐 Đang tải dữ liệu Google Sheet ở chế độ đọc Công Khai (Public CSV)...")
    csv_url = get_csv_url(url)
    req = urllib.request.Request(csv_url, headers={'User-Agent': 'Mozilla/5.0'})
    with urllib.request.urlopen(req) as response:
        csv_data = response.read().decode('utf-8')
    df = pd.read_csv(io.StringIO(csv_data))
    return df

def rule_based_score(description_text):
    desc = str(description_text).lower()
    
    # 1. CHECK JUNK / RÁC RULES (-50 POINTS)
    is_junk = False
    reasons = []
    
    # Unrealistic price
    if (("giá 1-2 tỷ" in desc or "1 tỷ" in desc or "2 tỷ" in desc) and "quận 1" in desc) or \
       ("trung tâm" in desc and "vài trăm triệu" in desc) or \
       ("thuê nguyên căn" in desc and "2 triệu" in desc) or \
       ("trung tâm thành phố" in desc and "2 triệu" in desc):
        is_junk = True
        reasons.append("Yêu cầu phi thực tế về giá cả")
        
    # No demand
    if "nhầm số" in desc or "không có nhu cầu" in desc or "dữ liệu cũ" in desc or "nhầm ngành" in desc:
        is_junk = True
        reasons.append("Khách hàng không có nhu cầu hoặc nhầm số/dữ liệu cũ")
        
    # No goodwill
    if "hỏi giá cho vui" in desc or "chưa có ý định mua" in desc or "thái độ không hợp tác" in desc:
        is_junk = True
        reasons.append("Khách hàng không thiện chí/hỏi giá cho vui")
        
    # Spam/Advertising
    if "spam" in desc or "quảng cáo" in desc or "bảo hiểm" in desc or "mời chào" in desc or "vay vốn" in desc:
        is_junk = True
        reasons.append("Nội dung spam hoặc quảng cáo ngược")
        
    # Contact error
    if "thuê bao" in desc or "không bắt máy" in desc or "không phản hồi" in desc or "gọi nhiều lần" in desc:
        is_junk = True
        reasons.append("Thông tin liên lạc lỗi (thuê bao/không nghe máy)")
        
    if is_junk:
        return -50, "Rác", " | ".join(reasons)
        
    # 2. CHECK VIP / SIÊU TIỀM NĂNG RULES (+50 POINTS)
    is_vip = False
    vip_reasons = []
    
    # Check if this is a mid-range case that should be overridden to Normal (0 points)
    is_mid_range = False
    if "đất nền vùng ven" in desc or "long an" in desc or "đồng nai" in desc or "2-3 tỷ" in desc:
        is_mid_range = True
    if "căn hộ 2pn" in desc or "quận 7" in desc or "4-5 tỷ" in desc or "gia đình trẻ" in desc:
        is_mid_range = True
    if "thuê mặt bằng" in desc and "dưới 50 triệu" in desc:
        is_mid_range = True
        
    if not is_mid_range:
        # Large budget
        if "tài chính cực mạnh" in desc or "tài chính mạnh" in desc or "không thành vấn đề" in desc or \
           "trên 30 tỷ" in desc or "thanh toán thẳng" in desc or "20 tỷ" in desc or "ngân sách lớn" in desc:
            is_vip = True
            vip_reasons.append("Tài chính lớn/cực mạnh")
            
        # Premium property type
        if "biệt thự đơn lập" in desc or "penthouse" in desc or "shophouse mặt đường lớn" in desc or \
           "quỹ đất công nghiệp" in desc or "sàn văn phòng diện tích lớn" in desc or "trên 2000m2" in desc:
            is_vip = True
            vip_reasons.append("Quan tâm sản phẩm bất động sản cao cấp")
            
        # Prime location
        if ("vinhomes ocean park" in desc or "phú mỹ hưng" in desc or "ven sông" in desc or "khu đông" in desc or "quận 1" in desc) and \
           ("vip" in desc or "biệt thự" in desc or "shophouse" in desc or "doanh nghiệp" in desc or "tài chính" in desc or "sỉ" in desc or "sàn văn phòng" in desc):
            is_vip = True
            vip_reasons.append("Vị trí đắc địa / Khu vực cao cấp")
            
        # Target client
        if "chủ doanh nghiệp" in desc or "nhà đầu tư chuyên nghiệp" in desc or "mua sỉ" in desc or "gom sỉ" in desc or "số lượng lớn" in desc:
            is_vip = True
            vip_reasons.append("Khách hàng là Chủ doanh nghiệp / Nhà đầu tư sỉ")
            
        # Urgency & Transparency
        if "pháp lý chuẩn 100%" in desc or "sổ hồng riêng" in desc or "gặp trực tiếp chủ đầu tư" in desc or "gặp trực tiếp giám đốc" in desc:
            is_vip = True
            vip_reasons.append("Yêu cầu pháp lý cao và giao dịch trực tiếp")
            
    if is_vip:
        return 50, "VIP", " | ".join(vip_reasons)
        
    # 3. DEFAULT: NORMAL / TIỀM NĂNG TRUNG BÌNH (0 POINTS)
    normal_reasons = []
    if "đất nền" in desc or "long an" in desc or "đồng nai" in desc:
        normal_reasons.append("Khách mua đất nền vùng ven tầm trung")
    elif "căn hộ" in desc or "chung cư" in desc:
        normal_reasons.append("Khách mua chung cư tầm trung")
    elif "thuê mặt bằng" in desc:
        normal_reasons.append("Khách thuê mặt bằng kinh doanh")
    else:
        normal_reasons.append("Khách hàng có nhu cầu bất động sản thông thường")
        
    if "vay ngân hàng" in desc or "vay" in desc:
        normal_reasons.append("cần hỗ trợ tài chính/vay ngân hàng")
        
    return 0, "Bình thường", " - ".join(normal_reasons)

def export_to_excel(df, filename):
    wb = Workbook()
    ws = wb.active
    ws.title = "Lead Scoring Reports"
    
    ws.views.sheetView[0].showGridLines = True
    
    headers = ["ID", "Tên Khách Hàng", "Số Điện Thoại", "Mô Tả Nhu Cầu", "Điểm Số", "Phân Loại", "Lý Do Chi Tiết"]
    ws.append(headers)
    
    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    thin_border_side = Side(border_style="thin", color="D3D3D3")
    cell_border = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)
    
    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = cell_border
        
    ws.row_dimensions[1].height = 28
    
    vip_fill = PatternFill(start_color="D1E7DD", end_color="D1E7DD", fill_type="solid")
    vip_font = Font(name="Segoe UI", size=10, color="0F5132", bold=True)
    
    normal_fill = PatternFill(start_color="E2F0D9", end_color="E2F0D9", fill_type="solid")
    normal_font = Font(name="Segoe UI", size=10, color="385723", bold=True)
    
    junk_fill = PatternFill(start_color="F8D7DA", end_color="F8D7DA", fill_type="solid")
    junk_font = Font(name="Segoe UI", size=10, color="842029", bold=True)
    
    for r_idx, row in enumerate(df.itertuples(index=False), start=2):
        ws.append([
            row.id,
            row.ten_khach,
            str(row.sdt),
            row.nhu_cau_mo_ta,
            row.diem,
            row.phan_loai,
            row.ly_do_chi_tiet
        ])
        ws.row_dimensions[r_idx].height = 24
        
        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=r_idx, column=col_idx)
            cell.border = cell_border
            cell.font = Font(name="Segoe UI", size=10)
            
            if col_idx in [1, 3, 5, 6]:
                cell.alignment = Alignment(horizontal="center", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
                
            if col_idx == 6:
                val = str(cell.value)
                if val == "VIP":
                    cell.fill = vip_fill
                    cell.font = vip_font
                elif val == "Rác":
                    cell.fill = junk_fill
                    cell.font = junk_font
                else:
                    cell.fill = normal_fill
                    cell.font = normal_font
                    
    column_widths = {
        1: 8,   # ID
        2: 22,  # Tên Khách Hàng
        3: 15,  # SĐT
        4: 55,  # Mô Tả Nhu Cầu
        5: 10,  # Điểm Số
        6: 15,  # Phân Loại
        7: 45   # Lý Do Chi Tiết
    }
    
    for col_idx, width in column_widths.items():
        col_letter = get_column_letter(col_idx)
        ws.column_dimensions[col_letter].width = width
        
    wb.save(filename)
    print(f"📄 Đã lưu báo cáo Excel tại: {filename}")

def main():
    sheet_url = "https://docs.google.com/spreadsheets/d/1joAy1H6PU19kwgsn57CSk_8cdci6vcDmME21CV_4n6E/edit?usp=sharing"
    
    df = load_sheet_data_secure(sheet_url)
    if df is None:
        print("❌ Lỗi khi tải dữ liệu. Kết thúc chương trình.")
        return

    # Check required columns
    required = ['id', 'ten_khach', 'sdt', 'nhu_cau_mo_ta']
    for col in required:
        if col not in df.columns:
            print(f"❌ Sheet thiếu cột bắt buộc: {col}")
            return
            
    print(f"📊 Đã tải thành công {len(df)} dòng dữ liệu.")
    print("⚡ Đang bắt đầu chấm điểm khách hàng dựa trên quy tắc nghiệp vụ offline...")
    
    scores = []
    classes = []
    reasons = []
    
    for idx, row in df.iterrows():
        diem, phan_loai, ly_do = rule_based_score(row['nhu_cau_mo_ta'])
        scores.append(diem)
        classes.append(phan_loai)
        reasons.append(ly_do)
        
    df['diem'] = scores
    df['phan_loai'] = classes
    df['ly_do_chi_tiet'] = reasons
    
    # Show stats
    total = len(df)
    vips = sum(1 for c in classes if c == "VIP")
    normals = sum(1 for c in classes if c == "Bình thường")
    junks = sum(1 for c in classes if c == "Rác")
    
    print("\n--- KẾT QUẢ THỐNG KÊ ---")
    print(f"Total Leads: {total}")
    print(f"VIP Leads:   {vips} ({vips/total*100:.1f}%)")
    print(f"Normal:      {normals} ({normals/total*100:.1f}%)")
    print(f"Junk Leads:  {junks} ({junks/total*100:.1f}%)")
    print("------------------------\n")
    
    # Export to Excel
    export_to_excel(df, "bao_cao_lead_scoring_bat_dong_san.xlsx")
    print("🎉 Hoàn thành chấm điểm thành công!")

if __name__ == "__main__":
    main()
